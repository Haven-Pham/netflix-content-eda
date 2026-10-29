"""Builds the Excel results workbook from the CSV tables in reports/tables.

    python scripts/build_excel_report.py

Run scripts/run_notebook.py first so the tables are current.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
TABLES = ROOT / "reports" / "tables"
FIGURES = ROOT / "reports" / "figures"

CONFIG = {'file': 'Netflix_Catalogue_EDA_Results.xlsx',
 'title': 'Netflix Catalogue Exploratory Data Analysis',
 'subtitle': 'What the 2021 catalogue contains, who it targets and how it grew',
 'findings': ['Films are 69.7% of the 8,797 titles, but the TV share of new additions rose from 25% in 2018 '
              'to 34% in 2021.',
              '67.3% of TV shows have a single season; the median film runs 98 minutes.',
              'TV-MA is the most common rating; adult-only ratings cover 45.5% of titles.',
              'Additions peaked in 2019 (2,016 titles); July and December receive the most new titles.',
              '46.6% of titles carry an International tag; India is the second-largest source country.'],
 'sheets': {'Key findings': ('key_findings.csv', 'headline metrics'),
            'Additions by year': ('titles_added_by_year.csv', 'films, series and TV share added per year'),
            'Ratings': ('ratings.csv', 'titles per audience rating'),
            'Genres': ('genres.csv', 'titles per genre tag'),
            'Countries': ('top_countries.csv', 'ten most frequent producing countries')},
 'figure': '04_additions_by_year.png',
 'source': 'Netflix Movies and TV Shows (Kaggle), collected September 2021'}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F4E79")


def number_format(column, series):
    name = column.lower()
    if not pd.api.types.is_numeric_dtype(series):
        return None
    if name.endswith("pct") or "_pct" in name or name.endswith("rate"):
        return '0.0"%"'
    if any(key in name for key in ("items", "orders", "sessions", "customers", "titles", "count")):
        return "#,##0"
    if any(key in name for key in ("usd", "revenue", "profit", "monetary", "refund", "margin", "aov", "per_session",
                                   "per_customer")):
        return '"$"#,##0.00' if series.abs().max() < 100 else '"$"#,##0'
    if pd.api.types.is_integer_dtype(series):
        return "#,##0"
    return "#,##0.00"


def add_table_sheet(wb, title, frame):
    ws = wb.create_sheet(title[:31])
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, column in enumerate(frame.columns, start=1):
        fmt = number_format(column, frame[column])
        letter = get_column_letter(idx)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        width = max(len(str(column)), *(len(str(v)) for v in frame[column].head(200))) if len(frame) else len(column)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = CONFIG["title"]
    ws["A1"].font = TITLE_FONT
    ws["A2"] = CONFIG["subtitle"]
    ws["A2"].font = Font(italic=True, color="555555")
    row = 4
    ws.cell(row, 1, "Key findings").font = Font(bold=True, size=12)
    for finding in CONFIG["findings"]:
        row += 1
        ws.cell(row, 1, f"• {finding}").alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
    row += 2
    ws.cell(row, 1, "Sheets in this workbook").font = Font(bold=True, size=12)
    for sheet, (_, description) in CONFIG["sheets"].items():
        row += 1
        ws.cell(row, 1, f"{sheet}: {description}")
    row += 2
    ws.cell(row, 1, "Source: " + CONFIG["source"]).font = Font(color="555555")
    ws.column_dimensions["A"].width = 120
    if CONFIG.get("figure") and (FIGURES / CONFIG["figure"]).exists():
        image = Image(str(FIGURES / CONFIG["figure"]))
        scale = 620 / image.width
        image.width, image.height = image.width * scale, image.height * scale
        ws.add_image(image, f"A{row + 2}")

    for sheet, (csv_name, _) in CONFIG["sheets"].items():
        add_table_sheet(wb, sheet, pd.read_csv(TABLES / csv_name))

    out = ROOT / "reports" / CONFIG["file"]
    wb.save(out)
    print(f"Wrote {out.relative_to(ROOT)} with {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    main()
